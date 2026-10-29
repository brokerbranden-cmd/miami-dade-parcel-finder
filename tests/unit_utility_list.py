"""Utility shutoff-list import slot: a CSV in $RAW/utility_shutoffs/<county>/ becomes the top signal 'Utility shutoff list' (bit 1, best code 1),
matched to parcels by service address (+ zip), for every county. Runs scripts/build_utility.py into a temp file (data/ is not touched)."""
import os, sys, json, gzip, subprocess, tempfile, struct, csv, datetime as dt
REPO = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'); errs = []
def ok(c, m):
    print(('PASS ' if c else 'FAIL ') + m, flush=True)
    if not c: errs.append(m)
def pack(path):
    b = gzip.decompress(open(path, 'rb').read()); hl = struct.unpack('<I', b[:4])[0]; h = json.loads(b[4:4 + hl]); base = 4 + hl; o = {}
    for c in h['cols']:
        if c['type'] == 'str' and c['name'] in ('folio', 'addr'): o[c['name']] = b[base + c['offset']: base + c['offset'] + c['bytes']].decode().split('\n')
    return o, h['n']
for cid, d in (('miami-dade', 'data/'), ('broward', 'data/broward/'), ('monroe', 'data/monroe/')):
    p, n = pack(os.path.join(REPO, d, 'main.bin.gz')); from collections import Counter
    cnt = Counter(p['addr'][:n]); pick = [(p['folio'][i], p['addr'][i]) for i in range(0, n, max(1, n // 400)) if p['addr'][i] and cnt[p['addr'][i]] == 1][:3]
    with tempfile.TemporaryDirectory() as raw:
        ud = os.path.join(raw, 'utility_shutoffs', cid); os.makedirs(ud)
        with open(os.path.join(ud, 'test_water_utility.csv'), 'w', newline='') as fh:
            w = csv.writer(fh); w.writerow(['Account', 'Service Address', 'Shutoff Date', 'Service Type'])
            day = (dt.date.today() - dt.timedelta(days=20)).strftime('%m/%d/%Y')
            for f, a in pick: w.writerow(['x', a.lower(), day, 'Water'])
            w.writerow(['x', '1 NOWHERE IMAGINARY ST', day, 'Water'])
        out = os.path.join(raw, 'u.json.gz'); r = subprocess.run([sys.executable, os.path.join(REPO, 'scripts', 'build_utility.py'), cid], env={**os.environ, 'RAW': raw, 'UT_OUT': out}, capture_output=True, text=True)
        ok(r.returncode == 0, f'{cid}: build ran ({r.stderr[-300:]})')
        if r.returncode: continue
        U = json.load(gzip.open(out)); R = U['rows']
        ok(all(f in R and R[f][0] & 1 and R[f][1] == 1 and R[f][3] == 'Water shut off' and R[f][4].startswith('Utility shutoff list') for f, _ in pick), f'{cid}: {len(pick)} listed addresses -> top signal Utility shutoff list')
        ok(U['counts']['list'] >= len(pick) and U['lists'][0]['rows'] == len(pick) + 1 and U['lists'][0]['matched'] == len(pick), f"{cid}: list stats {U['lists']}")
# open utility-lien listing (Margate FY2027-305 shape): title row, multi-line 'street \n\n CITY, FL  zip' cells, kind 'lien'
import openpyxl
p, n = pack(os.path.join(REPO, 'data/broward/main.bin.gz')); cnt = Counter(p['addr'][:n]); pick = [(p['folio'][i], p['addr'][i]) for i in range(7, n, max(1, n // 300)) if p['addr'][i] and cnt[p['addr'][i]] == 1][:2]
with tempfile.TemporaryDirectory() as raw:
    ud = os.path.join(raw, 'utility_shutoffs', 'broward'); os.makedirs(ud); wb = openpyxl.Workbook(); ws = wb.active
    ws.append(['Open Utility Liens', None]); ws.append(['Location Address Without Jurisdiction', 'Lien Open Date'])
    for k, (f, a) in enumerate(pick): ws.append([a + '  \n\nSOMEWHERE, FL' + ('  00000' if k else ''), dt.datetime.now() - dt.timedelta(days=30)]); ws.append([None, None])
    wb.save(os.path.join(ud, 'liens.xlsx'))
    json.dump({'source': 'Test open utility liens', 'kind': 'lien', 'header_row': 2, 'address_col': 'Location Address Without Jurisdiction', 'date_col': 'Lien Open Date'}, open(os.path.join(ud, 'liens.meta.json'), 'w'))
    out = os.path.join(raw, 'u.json.gz'); r = subprocess.run([sys.executable, os.path.join(REPO, 'scripts', 'build_utility.py'), 'broward'], env={**os.environ, 'RAW': raw, 'UT_OUT': out}, capture_output=True, text=True)
    ok(r.returncode == 0, f'lien: build ran ({r.stderr[-300:]})')
    if not r.returncode:
        U = json.load(gzip.open(out)); R = U['rows']; L = U['lists'][0]
        ok(L['rows'] == 2, f'lien: blank / title rows skipped {L}')
        ok(L['matched'] == 1 and pick[1][0] not in R, f'lien: zip from the address cell is enforced {L}')
        f0 = pick[0][0]; ok(f0 in R and R[f0][1] == 1 and R[f0][3] == 'Utility lien (open)' and R[f0][4] == 'Test open utility liens' and R[f0][2], f'lien: top signal row {R.get(f0)}')
# per-city parsers (scripts/utility_lists.py): Pompano Beach XLSX shape -> common table, personal columns never written, county lock
U_LISTS = os.path.join(REPO, 'scripts', 'utility_lists.py')
with tempfile.TemporaryDirectory() as raw:
    src = os.path.join(raw, 'src.xlsx'); wb = openpyxl.Workbook(); ws = wb.active
    ws.append(['Work Request Header Assign Date', 'Meter Location Street Address', 'Work Request Header Completion Code', 'Work Request Header Close Date', 'Customer Name', 'Account Number'])
    ws.append([dt.datetime(2026, 9, 1), '2511  NW 11 ST', 'PBF', dt.datetime(2026, 9, 9), 'JANE TESTPERSON', '000123456'])
    ws.append([dt.datetime(2026, 9, 2), None, 'FNLF', dt.datetime(2026, 9, 10), 'JOHN TESTPERSON', '000654321'])
    ws.append([dt.datetime(2026, 9, 3), '1020 N RIVERSIDE DR', 'FNLF', None, 'ANN TESTPERSON', '000111222']); ws.append([3, None, None, None, None, None])
    wb.save(src)
    r = subprocess.run([sys.executable, U_LISTS, 'broward', 'pompano_beach', 'pb_test', src], env={**os.environ, 'RAW': raw}, capture_output=True, text=True)
    ok(r.returncode == 0, f'parser: pompano_beach ran ({r.stderr[-300:]})')
    if not r.returncode:
        ud = os.path.join(raw, 'utility_shutoffs', 'broward'); rows = list(csv.DictReader(open(os.path.join(ud, 'pb_test.csv')))); M = json.load(open(os.path.join(ud, 'pb_test.meta.json')))
        ok(list(rows[0].keys()) == ['address', 'city', 'disconnect_date', 'service_type', 'status', 'source', 'source_file'], f'parser: common columns only {list(rows[0].keys())}')
        allt = open(os.path.join(ud, 'pb_test.csv')).read() + open(os.path.join(ud, 'unmatched', 'pb_test_parse_skipped.csv')).read() + json.dumps(M)
        ok('TESTPERSON' not in allt and '000123456' not in allt and '000654321' not in allt, 'parser: names / account numbers are not written anywhere')
        ok(len(rows) == 2 and M['skipped'] == 1 and rows[0]['address'] == '2511 NW 11 ST' and rows[0]['disconnect_date'] == '2026-09-09' and rows[1]['disconnect_date'] == '2026-09-03'
           and rows[0]['status'] == 'completion code PBF' and rows[0]['service_type'] == 'water', f'parser: rows / dates (close date, else assign date) / count row skipped {rows}')
        ok(M['address_col'] == 'address' and M['date_col'] == 'disconnect_date' and M['cities'][0] == 'Pompano Beach' and M['short'], f'parser: sidecar meta for the builder {M}')
    r = subprocess.run([sys.executable, U_LISTS, 'miami-dade', 'pompano_beach', 'x', src], env={**os.environ, 'RAW': raw}, capture_output=True, text=True)
    ok(r.returncode != 0 and 'never mixed' in (r.stderr + r.stdout), 'parser: a Broward city parser refuses another county')
# builder fallbacks for OCR'd / work-order addresses: no spaces, missing suffix, missing direction, Pompano abbreviations; service-area cities;
# condo unit left unmatched with a reason; list status kept as the 6th row field; short place name map
p, n = pack(os.path.join(REPO, 'data/broward/main.bin.gz')); cnt = Counter(p['addr'][:n])
def find(rx, k=0):
    import re as _re
    out = [(p['folio'][i], p['addr'][i]) for i in range(n) if p['addr'][i] and cnt[p['addr'][i]] == 1 and _re.fullmatch(rx, p['addr'][i])]
    return out[k] if len(out) > k else (None, None)
fa, aa = find(r'\d{3,4} (NW|NE|SW|SE) \d{1,2} (ST|AVE|TER)', 5); fb, ab = find(r'\d{3,4} (NW|NE|SW|SE) \d{1,2} (WAY|TER)', 9); fc, ac = find(r'\d{3,4} [A-Z]{4,} (WAY)', 2)
with tempfile.TemporaryDirectory() as raw:
    ud = os.path.join(raw, 'utility_shutoffs', 'broward'); os.makedirs(ud); day = (dt.date.today() - dt.timedelta(days=10)).isoformat()
    cases = [(aa.replace(' ', ''), 'compact'), (ab.rsplit(' ', 1)[0], 'prefix'), (ac.replace(' WAY', ' WY'), 'abbrev'), ('1 NOWHERE IMAGINARY ST', 'none')]
    with open(os.path.join(ud, 'ocr_list.csv'), 'w', newline='') as fh:
        w = csv.writer(fh); w.writerow(['address', 'disconnect_date', 'status'])
        for a, _ in cases: w.writerow([a, day, 'on cut list'])
    json.dump({'source': 'Test City water cut list', 'short': 'Test City', 'kind': 'water', 'ocr': True, 'address_col': 'address', 'date_col': 'disconnect_date'}, open(os.path.join(ud, 'ocr_list.meta.json'), 'w'))
    out = os.path.join(raw, 'u.json.gz'); r = subprocess.run([sys.executable, os.path.join(REPO, 'scripts', 'build_utility.py'), 'broward'], env={**os.environ, 'RAW': raw, 'UT_OUT': out}, capture_output=True, text=True)
    ok(r.returncode == 0, f'fallbacks: build ran ({r.stderr[-300:]})')
    if not r.returncode:
        U = json.load(gzip.open(out)); R = U['rows']; L = U['lists'][0]
        ok(all(f in R for f in (fa, fb, fc)), f'fallbacks: {cases[:3]} -> {[R.get(f) for f in (fa, fb, fc)]} rules {L.get("rules")}')
        ok(L['matched'] == 3 and L['parcels'] == 3 and L['short'] == 'Test City' and U['short'].get('Test City water cut list') == 'Test City', f'fallbacks: list stats {L}')
        ok(R.get(fa, [0] * 6)[5:] == ['on cut list'] and R[fa][2] == day and R[fa][3] == 'Water shut off', f'fallbacks: row carries date + status {R.get(fa)}')

# Dania Beach "Work Report - Post Shut-Off" (text-layer PDF, read by word x position): long names running into the address column,
# sprinkler-meter suffixes (S/S, SPSY, SP/SY, SPR/SYS mid-address), wrapped address words, repeats -> latest date, never names / accounts / phones
import importlib.util, re
spec = importlib.util.spec_from_file_location('ulists', U_LISTS); UL = importlib.util.module_from_spec(spec); spec.loader.exec_module(UL)
def _line(y, items): return [(x, y, x + 6 * len(t), t) for x, t in items]
def _acct(y, acct, name, addr, x_addr=214.9, wrap=None):
    w = [(36.0, acct)]; x = 98.7; nw = name.split(); step = min(6.0, (x_addr - 5 - 98.7 - 3 * len(nw)) / max(1, len(''.join(nw))))   # long name squeezed up to the address x (real layout)
    for t in nw: w.append((x, t)); x += step * len(t) + 3
    x = x_addr
    for t in addr.split(): w.append((x, t)); x += 6 * len(t) + 3
    w += [(390.0, 'Home'), (420.0, '(999)'), (450.0, '555-0100'), (520.0, 'Delinquent'), (600.0, '$12.34')]
    out = _line(y, w)
    if wrap: out += _line(y + 7, [(x_addr, wrap)])
    return out + _line(y + 13.7, [(x_addr, 'Dania'), (x_addr + 24, 'Beach,'), (x_addr + 50, 'FL,'), (x_addr + 62, '33004'), (390.0, 'Work')]) + _line(y + 30, [(50.0, '999000111'), (130.0, '5/8')])
def _page(date, blocks): return _line(40, [(400.0, 'Shut-Off'), (440.0, '-'), (446.0, date)]) + [w for b in blocks for w in b]
fake = {'a.pdf': [_page('09/01/2026', [_acct(100, '900001-001', 'JANE TESTPERSON', '101 SE 9 LN S/S'), _acct(150, '900002-001', 'LONGNAME TESTCO 161 LLC', '100 NE 1 AVE'),
                                     _acct(200, '900003-001', 'TEST EATERY OF X', '500 S FEDERAL SPR/SYS HWY'), _acct(250, '900003-002', 'TEST EATERY OF X', '500 S FEDERAL HWY')])],
        'b.pdf': [_page('09/29/2026', [_acct(100, '900001-002', 'JANE TESTPERSON', '101 SE 9 LN S/S'), _acct(150, '900004-001', 'JOHN TESTPERSON', '900 S DIXIE (AKA PHIPPEN RD)', wrap='HWY'),
                                     _acct(200, '900005-001', 'ANN TESTPERSON', '40 SW 19 ST WBILL')])]}
UL.bbox_pages = lambda f: fake[os.path.basename(f)]
rows, sk, M = UL.parse_dania_beach(['/x/a.pdf', '/x/b.pdf']); A = {(r['address'], r['status']): r for r in rows}
ok(len(rows) == 6 and not sk and M['repeats'] == 1 and M['accounts_total'] == 7, f'dania: 7 accounts -> 6 rows (1 repeat collapsed) {sorted(A)} skipped {sk} meta {M}')
ok(A.get(('101 SE 9 LN', 'water shut off for nonpayment (sprinkler meter)'), {}).get('disconnect_date') == '2026-09-29', f'dania: sprinkler suffix stripped, repeat keeps the latest date {rows}')
ok(('100 NE 1 AVE', 'water shut off for nonpayment') in A, f'dania: name running into the address column is cut at the address x ({sorted(A)})')
ok(('500 S FEDERAL HWY', 'water shut off for nonpayment (sprinkler meter)') in A and ('500 S FEDERAL HWY', 'water shut off for nonpayment') in A, 'dania: mid-address SPR/SYS = separate sprinkler-meter row')
ok(('900 S DIXIE HWY', 'water shut off for nonpayment') in A and ('40 SW 19 ST', 'water shut off for nonpayment') in A, 'dania: wrapped address word joined, AKA / WBILL notes dropped')
ok(all(r['city'] == 'Dania Beach' and r['service_type'] == 'water' and list(r) == UL.COLS for r in rows), 'dania: common columns, city + water')
allt = json.dumps(rows) + json.dumps(M)
ok(not re.search(r'TESTPERSON|TESTCO|TEST EATERY|9000\d\d-|555-0100|999000111|\(999\)|\$', allt), 'dania: no names / account / phone / meter numbers / balances in the output')
# Wilton Manors Off List shape (clean XLSX path): header row 1, blank row 2, status picked = the date column -> no date echo; 'FRONT' unit; meta undated 'older'
fw, aw = find(r'\d{3,4} (NW|NE|SW|SE) \d{1,2} (ST|AVE|TER)', 11); fx, ax = find(r'\d{3,4} (NW|NE|SW|SE) \d{1,2} (ST|AVE|TER)', 13)
with tempfile.TemporaryDirectory() as raw:
    ud = os.path.join(raw, 'utility_shutoffs', 'broward'); os.makedirs(ud); wb = openpyxl.Workbook(); ws = wb.active
    ws.append(['SERVICE ADDRESS', 'COMPLETION DATE', 'SERVICE']); ws.append([None, None, None])
    ws.append([aw + ' FRONT', dt.datetime.now() - dt.timedelta(days=20), 'WATER']); ws.append([ax, 'YEARS EMPTY LOT', 'WATER'])
    wb.save(os.path.join(ud, 'wm.xlsx')); json.dump({'source': 'Test WM water shutoff', 'kind': 'water', 'address_col': 'SERVICE ADDRESS', 'date_col': 'COMPLETION DATE', 'undated': 'older'}, open(os.path.join(ud, 'wm.meta.json'), 'w'))
    out = os.path.join(raw, 'u.json.gz'); r = subprocess.run([sys.executable, os.path.join(REPO, 'scripts', 'build_utility.py'), 'broward'], env={**os.environ, 'RAW': raw, 'UT_OUT': out}, capture_output=True, text=True)
    ok(r.returncode == 0, f'wilton shape: build ran ({r.stderr[-300:]})')
    if not r.returncode:
        U = json.load(gzip.open(out)); R = U['rows']; L = U['lists'][0]
        ok(L['rows'] == 2 and L['matched'] == 1 and L['older'] == 1 and L['rules'].get('base+unit') == 1, f'wilton shape: blank row 2 skipped, FRONT unit -> base address, undated note row = older {L}')
        ok(R.get(fw, [])[5:] == [] and R[fw][2], f'wilton shape: dated row has no status echo of the date {R.get(fw)}')
        ok(fx not in R, f'wilton shape: undated (years) row left out like rows older than 12 months {R.get(fx)}')
print('FAILED: ' + '; '.join(errs) if errs else 'ALL PASS'); sys.exit(1 if errs else 0)
